# promo_link_checker_v1.0.py
# 2026-09-29  Jonghyun Park w/ Claude
r"""
promo_link_checker — 국가별 홈 · 프로모션(offer) 페이지의 클릭 가능한 요소 중에
타겟 캠페인 링크가 걸려 있는지 매일 확인한다. 볼 페이지는 CHECK_PAGES 로 고른다.

전체 흐름:
  SITECODES × CHECK_PAGES → URL 생성 (home = /{sitecode}/, offer = /{sitecode}/offer/, hq 는 /hq/shop/)
  → Playwright 로 방문 · 쿠키 배너 닫기 · 끝까지 스크롤(레이지 로딩 링크 노출)
  → 모든 <a href> 의 URL 경로를 '/' 로 쪼개 각 구간을 KEYWORDS 규칙과 대조
  → 처음 보는 hit 요소(또는 디자인이 바뀐 요소)가 있을 때만 빨간 테두리를 그려 전체 페이지 캡처 (1회성)
  → output/_daily_report.xlsx 에 하루 1블록 누적 (최신이 D~F열, 과거는 오른쪽으로 밀림)

키워드 규칙 (KEYWORDS):
  한 줄 = 규칙 1개. 줄 안의 단어는 공백으로 구분하고, **URL 경로의 '/' 와 '/' 사이 한 구간 안에
  모두 들어 있어야(AND)** hit. 구간 안에서의 순서·사이 단어·구분자(- _ .)는 무관하다.
    예) "spring sale" → /spring-sale/ · /spring_sale/ · /spring-big-sale/ · /springsale/ hit
        /spring/summersales/ (두 구간에 나뉨) · ?campaign=spring-sale (쿼리) · 링크 텍스트만 일치 → 제외
  host(도메인)·쿼리 파라미터·#fragment 는 보지 않는다. 여러 줄이면 OR.

사용 예:
  python promo_link_checker_v1.0.py
  python promo_link_checker_v1.0.py --sitecodes ae,hq
  python promo_link_checker_v1.0.py --pages home
  python promo_link_checker_v1.0.py --sitecodes uk --debug --no-capture

출력 (스크립트 폴더 기준):
  output/_daily_report.xlsx                      일일 누적 리포트 (유일한 상시 산출물)
  output/capture/<MMDD>/<sitecode>_<page>_<HHMM>.png  새 요소·디자인 변경이 있을 때만
  output/_capture_state.json                     이미 캡처한 요소 기록 (지우면 처음부터 다시 캡처)
  output/_run_latest.log                         매 실행 덮어씀 (pythonw 스케줄 실행 추적용)

의존성: playwright, openpyxl   (설치 후 `python -m playwright install chromium` 1회)
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import logging
import os
import queue
import re
import shutil
import sys
import threading
import time
import zipfile
from datetime import datetime
from pathlib import Path
from urllib.parse import unquote, urlparse

from playwright.sync_api import sync_playwright, Page, TimeoutError as PlaywrightTimeoutError


# ════════════════════════════════════════════════════════════════════
# 사용자가 바꿔야 하는 부분
# ════════════════════════════════════════════════════════════════════

SCRIPT_DIR = Path(__file__).resolve().parent

# ─── 감지 키워드 ───────────────────────────────────────────────────
# 한 줄 = 규칙 1개 / 줄 안 단어는 AND / 줄끼리는 OR. 대소문자·구분자 무시, 부분일치.
# '#' 으로 시작하는 줄과 빈 줄은 무시한다. 따옴표 없이 그대로 적는다.
KEYWORDS = """
spring sale
"""

# ─── 대상 국가 ─────────────────────────────────────────────────────
# 한 줄에 sitecode 하나. 따옴표·쉼표 없이 적는다. '#' 줄·빈 줄은 무시.
# 2026-09-29 전수 점검(홈 92개 + /offer/): 87개 정상 + ru(/hero-products/) = 88개. 아래 '#' 3개는 offer 페이지 없음. (ge 는 운영 중단 sitecode 라 목록에서 제외)
# cn 은 도메인이 달라 SITE_OVERRIDES 로 https://www.company_name.com.cn/offer/ 를 연다.
SITECODES = """
ae
ae_ar
africa_en
africa_fr
africa_pt
al
ar
at
au
az
ba
bd
be
be_fr
bg
br
ca
ca_fr
ch
ch_fr
cl
cn
co
cz
de
dk
ee
eg
es
fi
fr
gr
hk
hk_en
hr
hu
id
ie
il
in
iq_ar
iq_ku
# iran    — /offer/ 없음 → 홈 리다이렉트, 대체 슬러그도 없음 (2026-09-29)
it
jp
kz_kz
kz_ru
latin
latin_en
lb
levant
levant_ar
lt
lv
mk
mm
# mn      — /offer/ 없음 → 홈 리다이렉트, 대체 슬러그도 없음 (2026-09-29)
mx
my
n_africa
nl
no
nz
pe
ph
pk
pl
# ps      — /offer/ 없음 → 홈 리다이렉트, 대체 슬러그도 없음 (2026-09-29)
pt
py
ro
rs
ru
sa
sa_en
se
hq
sg
si
sk
th
tr
tw
ua
uk
us
uy
uz_ru
uz_uz
vn
za
"""

# ─── 검사할 페이지 ─────────────────────────────────────────────────
# 한 줄에 하나, 따옴표 없이. 안 볼 페이지는 지우거나 앞에 '#'. (쓸 수 있는 값 = 아래 PAGE_SEGMENTS 의 키)
CHECK_PAGES = """
home
offer
"""

# ─── 대상 URL ──────────────────────────────────────────────────────
# URL = {BASE_DOMAIN}/{sitecode}/{페이지 슬러그}/   (슬러그가 "" 면 홈 = {BASE_DOMAIN}/{sitecode}/)
BASE_DOMAIN: str = "https://www.company_name.com"
PAGE_SEGMENTS: dict[str, str] = {
    "home": "",
    "offer": "offer",
}
# 기본 패턴을 안 따르는 사이트만:
#   base_domain / include_sitecode / segments({페이지: 그 사이트에서 쓸 슬러그})
SITE_OVERRIDES: dict[str, dict] = {
    "cn": {"base_domain": "https://www.company_name.com.cn", "include_sitecode": False},
    # 본사(hq)는 Offer 허브가 /hq/shop/ 이다 (/hq/offer/ 는 404).
    "hq": {"segments": {"offer": "shop"}},
    # ru 는 /ru/offer/ 가 홈으로 리다이렉트되고, 프로모션 허브가 /ru/hero-products/ 다.
    "ru": {"segments": {"offer": "hero-products"}},
}
# 리포트에 링크 설명으로 적을 속성 (보이는 텍스트가 없는 이미지 배너용). ⚠ 매칭에는 쓰지 않는다.
TEXT_ATTRIBUTES: list[str] = ["aria-label", "title"]
# "클릭 가능한 요소" 셀렉터
CLICKABLE_SELECTOR: str = "a[href], area[href]"
# True 면 GNB/Footer 안 링크는 hit 에서 뺀다 (False = 페이지 전체. 리포트엔 영역 태그가 붙는다)
EXCLUDE_GLOBAL_UI: bool = False
GLOBAL_UI_SELECTOR: str = "header, nav, footer, [class*='gnb'], [id*='gnb'], [an-ac='gnb']"

# ─── 출력 ──────────────────────────────────────────────────────────
OUTPUT_DIR: Path = SCRIPT_DIR / "output"
DAILY_REPORT_NAME: str = "_daily_report.xlsx"
CAPTURE_SUBDIR: str = "capture"
# hit 요소를 캡처한다. False 면 캡처를 아예 안 한다 (--no-capture 와 같음)
CAPTURE_ON_HIT: bool = True
CAPTURE_FULL_PAGE: bool = True
# ── 캡처 중복 방지 (1회성 캡처) ──
# 같은 요소(sitecode + page + 링크 URL)는 **한 번만** 캡처한다. 매일 hit 여도 다시 찍지 않는다.
# 단 그 요소의 디자인 지문(아래 FINGERPRINT_FIELDS)이 바뀌면 한 번 더 찍는다. 새 링크는 당연히 찍는다.
# 기록 파일을 지우면 처음부터 다시 캡처한다.
CAPTURE_STATE_NAME: str = "_capture_state.json"
# 디자인 지문에 넣을 항목:
#   size    = 요소 크기(FINGERPRINT_SIZE_STEP px 단위로 반올림)
#   classes = class 목록 (FINGERPRINT_CLASS_IGNORE 에 걸리는 상태 class 는 제외)
#   images  = 안쪽 <img> src · 배경 이미지 URL (쿼리 제외)
#   text    = 요소 텍스트 (숫자는 지운다 — 가격·날짜만 바뀐 건 디자인 변경으로 안 본다)
FINGERPRINT_FIELDS: list[str] = ["size", "classes", "images", "text"]
FINGERPRINT_SIZE_STEP: int = 20
# 캐러셀·레이지 로딩이 로드할 때마다 바꾸는 상태 class — 지문에서 뺀다 (정규식, 대소문자 무시)
FINGERPRINT_CLASS_IGNORE: str = r"active|current|visible|hidden|loaded|loading|lazy|hover|focus|prev|next|duplicate|selected|animat|is-|swiper-slide-"
# 캡처 시 테두리: 새로 찍는 이유가 된 요소(신규·디자인 변경) / 이미 찍은 적 있는 요소
HIGHLIGHT_STYLE: str = "4px solid red"
HIGHLIGHT_KNOWN_STYLE: str = "3px dashed orange"
# 리포트 셀에 적을 매칭 링크 최대 개수 (넘으면 "… 외 N건")
MAX_LINKS_IN_REPORT: int = 20
# 리포트 저장 전 최소 여유공간(MB). 0 이면 검사 안 함.
DAILY_REPORT_MIN_FREE_MB: int = 200
# 실행 중 건별 append 되는 안전망 CSV. 완주 후 삭제(KEEP_RESULT_CSV=False).
RESULT_CSV_NAME: str = "_result_latest.csv"
KEEP_RESULT_CSV: bool = False
# 실행 로그 (매 실행 덮어씀). pythonw 스케줄 실행은 콘솔이 없어서 이것만 남는다.
LOG_TO_FILE: bool = True
RUN_LOG_NAME: str = "_run_latest.log"

# ─── 브라우저 ──────────────────────────────────────────────────────
MAX_WORKERS: int = 4
BROWSER_HEADLESS: bool = True
BROWSER_CHANNEL: str | None = "chrome"     # None 이면 Playwright 번들 Chromium
BROWSER_TIMEOUT_MS: int = 60_000
BROWSER_WAIT_UNTIL: str = "load"
BROWSER_LOCALE: str = "en-US"
BROWSER_USER_AGENT: str | None = (
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
    "(KHTML, like Gecko) Chrome/135.0.0.0 Safari/537.36"
)
BROWSER_VIEWPORT: dict = {"width": 1440, "height": 900}
COOKIE_CONSENT_SELECTOR: str = "#truste-consent-button"
COOKIE_BANNER_WAIT_MS: int = 3_000
COOKIE_RELOAD_TIMEOUT_MS: int = 3_000
# 레이지 로딩 링크를 띄우기 위한 스크롤 (한 화면씩 내리며 잠깐 대기)
SCROLL_PAUSE_MS: int = 400
SCROLL_MAX_STEPS: int = 60
# 스크롤 후 추가 요청이 잦아들기를 기다리는 상한
SETTLE_NETWORKIDLE_MS: int = 8_000
# 접근 실패(타임아웃·브라우저 오류) 재시도 횟수. HTTP 4xx/5xx·리다이렉트는 재시도 안 함.
RETRY_COUNT: int = 1


# ════════════════════════════════════════════════════════════════════
# 내부 사용 — 보통 수정 불필요
# ════════════════════════════════════════════════════════════════════

LOGGER_NAME = "promo_link_checker"
logger = logging.getLogger(LOGGER_NAME)

RESULT_COLUMNS = ["sitecode", "page", "url", "final_url", "result", "keywords", "links",
                  "capture", "prev_capture", "new_elements", "detail", "checked_at"]
_RE_CSS_URL = re.compile(r"""url\(["']?([^"')]+)["']?\)""")
_RE_DIGITS = re.compile(r"\d+")
# 리포트: A=sitecode, B=page, C=url 고정, 이후 하루 3열 [결과, 매칭키워드, 매칭링크]
REPORT_FIXED_COLS = 3
REPORT_DAY_HEADERS = ("결과", "매칭키워드", "매칭링크")
RESULT_HIT, RESULT_MISS = "O", "X"
RESULT_FAIL_PREFIX = "접근실패"
RESULT_NOT_RUN = "(미실행)"

_RE_NON_WORD = re.compile(r"[\W_]+", re.UNICODE)

_thread_local = threading.local()

# 페이지 안에서 링크를 모으는 JS. 각 요소에 data-plc-idx 를 달아 캡처 때 다시 찾는다.
_JS_COLLECT = r"""({selector, textAttrs, uiSelector}) => {
    const out = [];
    document.querySelectorAll(selector).forEach((el, i) => {
        el.setAttribute("data-plc-idx", String(i));
        const text = (el.innerText || el.textContent || "").replace(/\s+/g, " ").trim();
        const extra = textAttrs.map(a => el.getAttribute(a) || "")
            .concat(Array.from(el.querySelectorAll("img[alt]")).map(img => img.getAttribute("alt") || ""))
            .filter(Boolean).join(" ");
        out.push({
            idx: i,
            href: el.getAttribute("href") || "",
            abs: el.href || "",
            text: text.slice(0, 300),
            extra: extra.slice(0, 300),
            global_ui: Boolean(uiSelector && el.closest(uiSelector)),
        });
    });
    return out;
}"""

# hit 요소의 디자인 지문 재료 (크기·class·이미지·텍스트)
_JS_DESCRIBE = r"""(idxs) => idxs.map(i => {
    const el = document.querySelector(`[data-plc-idx="${i}"]`);
    if (!el) return null;
    const r = el.getBoundingClientRect();
    const imgs = Array.from(el.querySelectorAll("img"))
        .map(im => im.currentSrc || im.getAttribute("src") || im.getAttribute("data-src") || "");
    const bgs = [el, ...Array.from(el.querySelectorAll("*")).slice(0, 50)]
        .map(n => getComputedStyle(n).backgroundImage).filter(b => b && b !== "none");
    return {
        idx: i, w: r.width, h: r.height,
        classes: el.getAttribute("class") || "",
        imgs: imgs, bgs: bgs,
        text: (el.innerText || el.textContent || "").replace(/\s+/g, " ").trim().slice(0, 300),
    };
})"""

_JS_HIGHLIGHT = r"""({idxs, style}) => {
    for (const i of idxs) {
        const el = document.querySelector(`[data-plc-idx="${i}"]`);
        if (el) { el.style.outline = style; el.style.outlineOffset = "2px"; }
    }
}"""


# ─── 설정 파싱 ─────────────────────────────────────────────────────

def _parse_lines(block: str) -> list[str]:
    """여러 줄 문자열 → 줄 목록 ('#' 줄·빈 줄 제외, 앞뒤 공백 제거)."""
    lines = []
    for line in block.splitlines():
        s = line.strip()
        if s and not s.startswith("#"):
            lines.append(s)
    return lines


def normalize_text(value: str) -> str:
    """소문자화 + 구분자(- _ / . 공백 등)를 공백 하나로. 매칭은 이 결과 위에서 부분일치로 한다."""
    return " " + _RE_NON_WORD.sub(" ", (value or "").lower()).strip() + " "


def parse_keyword_rules(block: str | None = None) -> list[tuple[str, list[str]]]:
    """KEYWORDS → [(원문 줄, [정규화 단어...]), ...]  (block 미지정 시 호출 시점의 KEYWORDS)"""
    rules = []
    for line in _parse_lines(KEYWORDS if block is None else block):
        words = [w for w in normalize_text(line).split() if w]
        if words:
            rules.append((line, words))
    return rules


def parse_sitecodes(block: str | None = None) -> list[str]:
    return list(dict.fromkeys(s.lower() for s in _parse_lines(SITECODES if block is None else block)))


def parse_pages(block: str | None = None) -> list[str]:
    pages = list(dict.fromkeys(s.lower() for s in _parse_lines(CHECK_PAGES if block is None else block)))
    unknown = [p for p in pages if p not in PAGE_SEGMENTS]
    if unknown:
        raise ValueError(f"CHECK_PAGES 에 모르는 페이지: {unknown} (가능: {list(PAGE_SEGMENTS)})")
    return pages


def path_segments(url: str) -> list[str]:
    """URL 경로를 '/' 로 쪼갠 구간들 (정규화). host·쿼리·#fragment 는 버린다."""
    path = urlparse(url or "").path
    return [normalize_text(unquote(seg)) for seg in path.split("/") if seg.strip()]


def match_rules(url: str, rules: list[tuple[str, list[str]]]) -> list[str]:
    """경로 구간 **하나**에 규칙 단어가 모두 들어 있으면 그 규칙 원문을 돌려준다.
    ⚠ 구간을 합쳐서 보면 /spring/summersales/ 처럼 서로 다른 구간의 단어가 섞여 걸린다."""
    segments = path_segments(url)
    return [line for line, words in rules
            if any(all(w in seg for w in words) for seg in segments)]


# ─── URL ───────────────────────────────────────────────────────────

def build_url(sitecode: str, page: str) -> str:
    override = SITE_OVERRIDES.get(sitecode, {})
    base = override.get("base_domain", BASE_DOMAIN).rstrip("/")
    segment = override.get("segments", {}).get(page, PAGE_SEGMENTS[page])
    parts = ([sitecode] if override.get("include_sitecode", True) else []) + ([segment] if segment else [])
    return "/".join([base] + parts) + "/"


def _same_page(requested: str, final: str) -> bool:
    """host(www 무시) + path(끝 슬래시 무시)가 같으면 같은 페이지로 본다 (쿼리·#fragment 무시)."""
    a, b = urlparse(requested), urlparse(final)
    host = lambda h: (h or "").lower().removeprefix("www.")
    return host(a.netloc) == host(b.netloc) and a.path.rstrip("/") == b.path.rstrip("/")


# ─── 브라우저 세션 ─────────────────────────────────────────────────

class BrowserSession:
    """Chromium 1개를 띄워 여러 페이지에 재사용한다. 스레드당 1개."""

    def __init__(self) -> None:
        self._playwright = None
        self._browser = None
        self.context = None

    def start(self) -> None:
        self._playwright = sync_playwright().start()
        kwargs = {"headless": BROWSER_HEADLESS}
        if BROWSER_CHANNEL:
            kwargs["channel"] = BROWSER_CHANNEL
        self._browser = self._playwright.chromium.launch(**kwargs)
        self._new_context()

    def _new_context(self) -> None:
        kwargs = {"locale": BROWSER_LOCALE, "viewport": BROWSER_VIEWPORT}
        if BROWSER_USER_AGENT:
            kwargs["user_agent"] = BROWSER_USER_AGENT
        self.context = self._browser.new_context(**kwargs)
        self.context.set_default_timeout(BROWSER_TIMEOUT_MS)

    def new_page(self) -> Page:
        try:
            return self.context.new_page()
        except Exception as e:
            # context 가 죽은 경우가 있어 한 번만 되살린다.
            logger.debug("new_page failed, recreating context: %s", e)
            try:
                self.context.close()
            except Exception:
                pass
            self._new_context()
            return self.context.new_page()

    def close(self) -> None:
        for obj in (self.context, self._browser):
            try:
                if obj:
                    obj.close()
            except Exception:
                pass
        try:
            if self._playwright:
                self._playwright.stop()
        except Exception:
            pass


def _get_thread_session() -> BrowserSession:
    session = getattr(_thread_local, "session", None)
    if session is None:
        session = BrowserSession()
        _thread_local.session = session       # start 가 중간에 실패해도 close 대상에 남도록 먼저 등록
        session.start()
    return session


def _close_thread_session() -> None:
    """현재 스레드의 세션을 닫는다 — 반드시 세션을 만든 워커 스레드 안에서 부를 것."""
    session = getattr(_thread_local, "session", None)
    if session is not None:
        session.close()
        _thread_local.session = None


def _dismiss_cookie_popup(page: Page, sitecode: str) -> None:
    try:
        page.wait_for_selector(COOKIE_CONSENT_SELECTOR, timeout=COOKIE_BANNER_WAIT_MS)
    except Exception:
        return
    # 동의 클릭이 같은 URL 재로드를 일으킬 수 있어 클릭 **전에** 네비게이션 대기를 건다.
    try:
        with page.expect_navigation(wait_until="load", timeout=COOKIE_RELOAD_TIMEOUT_MS):
            page.click(COOKIE_CONSENT_SELECTOR)
    except PlaywrightTimeoutError:
        pass
    except Exception as e:
        logger.debug("[%s] cookie click failed: %s", sitecode, e)


def _scroll_through(page: Page) -> None:
    """한 화면씩 끝까지 내려 레이지 로딩 영역(배너·카드)을 띄운 뒤 맨 위로 돌아온다."""
    step = BROWSER_VIEWPORT["height"]
    y = 0
    for _ in range(SCROLL_MAX_STEPS):
        height = page.evaluate("() => document.documentElement.scrollHeight")
        if y >= height:
            break
        y += step
        page.evaluate("y => window.scrollTo(0, y)", y)
        page.wait_for_timeout(SCROLL_PAUSE_MS)
    else:
        logger.debug("scroll step limit reached (%d)", SCROLL_MAX_STEPS)
    try:
        page.wait_for_load_state("networkidle", timeout=SETTLE_NETWORKIDLE_MS)
    except Exception:
        pass
    page.evaluate("() => window.scrollTo(0, 0)")


# ─── 캡처 중복 방지 (요소 지문 · 상태 파일) ─────────────────────────

def element_key(url: str) -> str:
    """요소 식별용 링크 키 = host + path (쿼리·#fragment·끝 슬래시 제외, 소문자)."""
    p = urlparse(url or "")
    return (p.netloc.lower().removeprefix("www.") + p.path.rstrip("/")).lower()


def _clean_src(src: str) -> str:
    if not src or src.startswith("data:"):
        return ""                      # 레이지 로딩 placeholder
    p = urlparse(src)
    return (p.netloc + p.path).lower()


def design_fingerprint(desc: dict) -> tuple[str, dict]:
    """요소 디자인 지문 (12자리 해시) + 사람이 볼 요약."""
    parts: dict = {}
    if "size" in FINGERPRINT_FIELDS:
        step = max(1, FINGERPRINT_SIZE_STEP)
        parts["size"] = [round(desc.get("w", 0) / step) * step, round(desc.get("h", 0) / step) * step]
    if "classes" in FINGERPRINT_FIELDS:
        ignore = re.compile(FINGERPRINT_CLASS_IGNORE, re.IGNORECASE) if FINGERPRINT_CLASS_IGNORE else None
        parts["classes"] = sorted({c for c in (desc.get("classes") or "").split()
                                   if not (ignore and ignore.search(c))})
    if "images" in FINGERPRINT_FIELDS:
        srcs = list(desc.get("imgs") or [])
        for bg in desc.get("bgs") or []:
            srcs += _RE_CSS_URL.findall(bg)
        parts["images"] = sorted({c for c in map(_clean_src, srcs) if c})
    if "text" in FINGERPRINT_FIELDS:
        parts["text"] = " ".join(_RE_DIGITS.sub("", desc.get("text") or "").lower().split())[:200]
    raw = json.dumps(parts, sort_keys=True, ensure_ascii=False)
    return hashlib.sha1(raw.encode("utf-8")).hexdigest()[:12], parts


def load_capture_state() -> dict:
    """{"sitecode|page": {링크키: {지문: {"first_seen", "capture", "design"}}}}"""
    path = OUTPUT_DIR / CAPTURE_STATE_NAME
    if not path.exists():
        return {}
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except Exception as e:
        # 깨진 상태 파일 때문에 매일 전부 재캡처되지 않도록 옆으로 치우고 경고만 남긴다
        broken = path.with_name(f"{path.name}.broken_{time.strftime('%y%m%d_%H%M')}")
        try:
            os.replace(path, broken)
        except OSError:
            pass
        logger.warning("⚠️ 캡처 기록 파일을 읽지 못해 새로 시작합니다: %s (%s)", broken.name, e)
        return {}


def save_capture_state(state: dict) -> None:
    path = OUTPUT_DIR / CAPTURE_STATE_NAME
    tmp = path.with_name(path.name + ".tmp")
    try:
        tmp.write_text(json.dumps(state, ensure_ascii=False, indent=1), encoding="utf-8")
        os.replace(tmp, path)
    except Exception as e:
        logger.warning("⚠️ 캡처 기록 저장 실패: %s", e)


def _state_key(sitecode: str, page_key: str) -> str:
    return f"{sitecode}|{page_key}"


# ─── 한 사이트 검사 ────────────────────────────────────────────────

def _result(sitecode, page, url, result, final_url="", keywords="", links="", capture="", detail="",
            prev_capture="", new_elements=""):
    return {
        "sitecode": sitecode, "page": page, "url": url, "final_url": final_url, "result": result,
        "keywords": keywords, "links": links, "capture": capture, "prev_capture": prev_capture,
        "new_elements": new_elements, "detail": detail,
        "checked_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
    }


def _format_link(item: dict) -> str:
    area = " [GNB/Footer]" if item["global_ui"] else ""
    label = item["text"] or item["extra"]
    label = f"  ({label[:50]})" if label else ""
    return f"{item['abs'] or item['href']}{label}{area}"


def check_site(sitecode: str, page_key: str, rules, capture: bool, run_date: str,
               known: dict | None = None) -> dict:
    """known = 이 (sitecode, page) 의 캡처 기록 {링크키: {지문: info}} — 읽기 전용."""
    known = known or {}
    url = build_url(sitecode, page_key)
    tag = f"{sitecode}/{page_key}"
    session = _get_thread_session()
    page = session.new_page()
    try:
        response = page.goto(url, wait_until=BROWSER_WAIT_UNTIL)
        status = response.status if response else None
        if status is not None and status >= 400:
            return _result(sitecode, page_key, url, f"{RESULT_FAIL_PREFIX}(HTTP {status})", page.url)
        if not _same_page(url, page.url):
            return _result(sitecode, page_key, url, f"{RESULT_FAIL_PREFIX}(redirect)", page.url,
                           detail=f"redirect → {page.url}")

        _dismiss_cookie_popup(page, tag)
        _scroll_through(page)

        items = page.evaluate(_JS_COLLECT, {
            "selector": CLICKABLE_SELECTOR,
            "textAttrs": TEXT_ATTRIBUTES,
            "uiSelector": GLOBAL_UI_SELECTOR,
        })
        logger.debug("[%s] clickable elements: %d", tag, len(items))

        matched_rules: list[str] = []
        hit_items: list[dict] = []
        for item in items:
            if EXCLUDE_GLOBAL_UI and item["global_ui"]:
                continue
            # 상대경로도 브라우저가 해석한 절대 URL(abs)로 본다
            hits = match_rules(item["abs"] or item["href"], rules)
            if not hits:
                continue
            hit_items.append(item)
            for h in hits:
                if h not in matched_rules:
                    matched_rules.append(h)

        if not hit_items:
            return _result(sitecode, page_key, url, RESULT_MISS, page.url,
                           detail=f"clickable {len(items)}")

        # ── 요소별 디자인 지문 → 이미 캡처한 요소인지 판정 ──
        descs = page.evaluate(_JS_DESCRIBE, [it["idx"] for it in hit_items])
        new_elems: dict[tuple[str, str], dict] = {}     # (링크키, 지문) → 신규/변경 정보
        new_idxs, known_idxs = [], []
        link_kind: dict[str, str] = {}                  # 링크키 → "new" / "changed" / "known"
        prev_captures: list[tuple[str, str]] = []       # (first_seen, capture)
        for it, desc in zip(hit_items, descs):
            ekey = element_key(it["abs"] or it["href"])
            fp, design = design_fingerprint(desc or {})
            seen_fps = known.get(ekey, {})
            if fp in seen_fps:
                known_idxs.append(it["idx"])
                info = seen_fps[fp]
                prev_captures.append((info.get("first_seen", ""), info.get("capture", "")))
                link_kind.setdefault(ekey, "known")
                continue
            kind = "changed" if seen_fps else "new"
            new_idxs.append(it["idx"])
            new_elems.setdefault((ekey, fp), {"key": ekey, "fp": fp, "kind": kind, "design": design})
            if link_kind.get(ekey) != "changed":
                link_kind[ekey] = kind

        # ── 리포트용 링크 목록 (링크당 1줄 + 신규/변경 태그) ──
        matched_links: list[str] = []
        seen_links: set[str] = set()
        for it in hit_items:
            ekey = element_key(it["abs"] or it["href"])
            if ekey in seen_links:
                continue
            seen_links.add(ekey)
            tag_txt = {"new": "  🆕신규", "changed": "  🔄디자인변경"}.get(link_kind.get(ekey, ""), "")
            matched_links.append(_format_link(it) + tag_txt)

        capture_rel = ""
        if capture and new_elems:
            try:
                page.evaluate(_JS_HIGHLIGHT, {"idxs": known_idxs, "style": HIGHLIGHT_KNOWN_STYLE})
                page.evaluate(_JS_HIGHLIGHT, {"idxs": new_idxs, "style": HIGHLIGHT_STYLE})
                rel = Path(CAPTURE_SUBDIR) / run_date / f"{sitecode}_{page_key}_{datetime.now():%H%M}.png"
                out = OUTPUT_DIR / rel
                out.parent.mkdir(parents=True, exist_ok=True)
                page.screenshot(path=str(out), full_page=CAPTURE_FULL_PAGE)
                capture_rel = rel.as_posix()
            except Exception as e:
                logger.warning("⚠️ %s: 캡처 실패 — %s", tag, e)

        prev = max((c for c in prev_captures if c[1]), default=("", ""))[1]
        return _result(sitecode, page_key, url, RESULT_HIT, page.url,
                       keywords=", ".join(matched_rules),
                       links="\n".join(matched_links),
                       capture=capture_rel,
                       prev_capture=prev,
                       # 캡처에 성공했을 때만 기록 대상으로 넘긴다 (못 찍었으면 다음 실행에서 다시 시도)
                       new_elements=json.dumps(list(new_elems.values()), ensure_ascii=False) if capture_rel else "",
                       detail=f"clickable {len(items)} / hit {len(hit_items)} / new {len(new_elems)}")
    finally:
        try:
            page.close()
        except Exception:
            pass


def check_site_with_retry(sitecode: str, page_key: str, rules, capture: bool, run_date: str,
                          known: dict | None = None) -> dict:
    last_error = ""
    for attempt in range(RETRY_COUNT + 1):
        try:
            return check_site(sitecode, page_key, rules, capture, run_date, known)
        except Exception as e:
            last_error = f"{type(e).__name__}: {str(e).splitlines()[0][:150]}"
            logger.debug("[%s/%s] attempt %d failed: %s", sitecode, page_key, attempt + 1, last_error)
    kind = "timeout" if "Timeout" in last_error else "error"
    return _result(sitecode, page_key, build_url(sitecode, page_key), f"{RESULT_FAIL_PREFIX}({kind})",
                   detail=last_error)


# ─── 결과 CSV (안전망) ──────────────────────────────────────────────

def _append_result_csv(path: Path, row: dict) -> None:
    new = not path.exists()
    with open(path, "a", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=RESULT_COLUMNS)
        if new:
            w.writeheader()
        w.writerow(row)


# ─── 일일 리포트 ───────────────────────────────────────────────────
# 일일 리포트 구조: 최신 날짜 블록이 항상 고정열 바로 오른쪽,
# 과거는 오른쪽으로 밀림 / 같은 날 재실행은 그 블록 재사용 / 대상 밖 행 보존 / 원자적 저장.

def write_daily_report(csv_path: Path, run_date: str, target_keys: list[tuple[str, str]]) -> Path | None:
    """target_keys = 이번 실행 대상 [(sitecode, page), ...] — 이 밖의 행은 손대지 않는다."""
    try:
        from openpyxl import Workbook, load_workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except Exception as e:
        logger.warning("⚠️ 일일 리포트 건너뜀(openpyxl 없음): %s", e)
        return None

    path = OUTPUT_DIR / DAILY_REPORT_NAME
    N = len(REPORT_DAY_HEADERS)
    first_day_col = REPORT_FIXED_COLS + 1

    try:
        with open(csv_path, encoding="utf-8-sig", newline="") as f:
            rows = list(csv.DictReader(f))
    except FileNotFoundError:
        logger.warning("⚠️ 일일 리포트 건너뜀(결과 CSV 없음): %s", csv_path)
        return None
    day_vals = {(r["sitecode"], r["page"]): r for r in rows if r.get("sitecode")}   # 뒤엣것이 최종
    targets = set(target_keys)

    if DAILY_REPORT_MIN_FREE_MB:
        try:
            free_mb = shutil.disk_usage(OUTPUT_DIR).free / (1024 * 1024)
            if free_mb < DAILY_REPORT_MIN_FREE_MB:
                logger.warning("⚠️ 일일 리포트 건너뜀(디스크 여유 %.0fMB) — 기존 파일 보존", free_mb)
                return None
        except OSError:
            pass

    if path.exists():
        try:
            wb = load_workbook(path)
        except Exception as e:
            broken = path.with_name(f"{path.name}.broken_{time.strftime('%y%m%d_%H%M')}")
            try:
                os.replace(path, broken)
                logger.warning("⚠️ 기존 리포트가 손상돼 치워두고 새로 만든다: %s (%s)", broken.name, e)
            except OSError as e2:
                logger.warning("⚠️ 일일 리포트 건너뜀(손상 파일 이동 실패): %s", e2)
                return None
            return write_daily_report(csv_path, run_date, target_keys)
        ws = wb.active
        # ⚠ openpyxl 의 insert_cols 는 병합 범위를 옮기지 않는다 → 헤더 병합을 풀고 끝에서 다시 건다.
        for rng in list(ws.merged_cells.ranges):
            ws.unmerge_cells(str(rng))
        existing = {ws.cell(row=1, column=c).value: c
                    for c in range(first_day_col, ws.max_column + 1, N)}
        if run_date in existing:
            base = existing[run_date]
        else:
            ws.insert_cols(first_day_col, N)
            base = first_day_col
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = "daily"
        base = first_day_col

    for c, name in enumerate(("sitecode", "page", "url"), start=1):
        ws.cell(row=1, column=c, value=name)
        ws.cell(row=2, column=c, value=name)
    ws.cell(row=1, column=base, value=run_date)
    for off, name in enumerate(REPORT_DAY_HEADERS):
        ws.cell(row=2, column=base + off, value=name)

    row_of = {}
    for r in range(3, ws.max_row + 1):
        sc, pg = ws.cell(row=r, column=1).value, ws.cell(row=r, column=2).value
        if sc:
            row_of[(str(sc), str(pg or ""))] = r
    next_row = max(ws.max_row + 1, 3)
    # 새 행은 실행 대상 순서(sitecode → page)대로 붙인다 — 병렬 완료 순서로 뒤섞이지 않게
    for key in list(target_keys) + [k for k in day_vals if k not in targets]:
        if key not in row_of:
            ws.cell(row=next_row, column=1, value=key[0])
            ws.cell(row=next_row, column=2, value=key[1])
            row_of[key] = next_row
            next_row += 1

    fill_hit = PatternFill("solid", fgColor="C6EFCE")
    fill_fail = PatternFill("solid", fgColor="FFC7CE")
    fill_none = PatternFill(fill_type=None)
    wrap = Alignment(wrap_text=True, vertical="top")
    top = Alignment(vertical="top")
    n_hit = n_fail = 0
    for key, r in row_of.items():
        if key in day_vals:
            v = day_vals[key]
            result, kw = v["result"], v["keywords"]
            links = [l for l in (v["links"] or "").split("\n") if l]
            if len(links) > MAX_LINKS_IN_REPORT:
                links = links[:MAX_LINKS_IN_REPORT] + [f"… 외 {len(links) - MAX_LINKS_IN_REPORT}건"]
            if v.get("capture"):
                links.append(f"📷 새 캡처: {v['capture']}")
            elif v.get("prev_capture"):
                links.append(f"📷 기존 캡처: {v['prev_capture']}")
            if result.startswith(RESULT_FAIL_PREFIX) and v.get("detail"):
                links.append(v["detail"])
            link_text = "\n".join(links)[:32000]
            ws.cell(row=r, column=3, value=v["url"])
        elif key not in targets:
            continue                    # 이번 실행 대상이 아니다 — 기존 칸 보존
        else:
            result, kw, link_text = RESULT_NOT_RUN, "", ""
            v = {}

        c_res = ws.cell(row=r, column=base, value=result)
        c_kw = ws.cell(row=r, column=base + 1, value=kw)
        c_ln = ws.cell(row=r, column=base + 2, value=link_text)
        c_res.alignment, c_kw.alignment, c_ln.alignment = top, wrap, wrap
        c_res.hyperlink = None
        if result == RESULT_HIT:
            n_hit += 1
            for c in (c_res, c_kw):
                c.fill = fill_hit
            snap = v.get("capture") or v.get("prev_capture")
            if snap:
                c_res.hyperlink = snap              # 리포트 기준 상대경로 → 클릭하면 캡처가 열린다
                c_res.font = Font(bold=True, underline="single", color="006100")
            else:
                c_res.font = Font(bold=True, color="006100")
        elif result.startswith(RESULT_FAIL_PREFIX) or result == RESULT_NOT_RUN:
            n_fail += 1
            c_res.fill, c_res.font = fill_fail, Font()
            c_kw.fill = fill_none
        else:
            c_res.fill, c_kw.fill, c_res.font = fill_none, fill_none, Font()

    # 헤더 병합 복구 + 서식
    for c in range(first_day_col, ws.max_column + 1, N):
        if ws.cell(row=1, column=c).value:
            ws.merge_cells(start_row=1, start_column=c, end_row=1, end_column=c + N - 1)
            ws.cell(row=1, column=c).alignment = Alignment(horizontal="center")
            ws.cell(row=1, column=c).font = Font(bold=True)
            for off in range(N):
                ws.cell(row=2, column=c + off).font = Font(bold=True)
                ws.column_dimensions[get_column_letter(c + off)].width = (10, 22, 70)[off]
    for c in range(1, REPORT_FIXED_COLS + 1):
        ws.cell(row=1, column=c).font = Font(bold=True)
        ws.cell(row=2, column=c).font = Font(bold=True)
    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 8
    ws.column_dimensions["C"].width = 44
    ws.freeze_panes = ws.cell(row=3, column=first_day_col)
    # ⚠ openpyxl 의 insert_cols 는 셀 값은 옮기지만 하이퍼링크의 ref(좌표)는 그대로 둔다.
    #   안 맞추면 과거 블록의 캡처 링크가 오늘 칸(D열)에 겹쳐 저장돼 엉뚱한 캡처가 열린다.
    for row_cells in ws.iter_rows(min_row=3):
        for cell in row_cells:
            if cell.hyperlink is not None:
                cell.hyperlink.ref = cell.coordinate

    # 저장: 임시파일 → 필수 파트 확인 → 원자적 교체 (저장 도중 실패해도 누적 이력이 안 날아가게)
    tmp = path.with_name(path.name + ".tmp")
    try:
        wb.save(tmp)
        with zipfile.ZipFile(tmp) as zf:
            missing = {"[Content_Types].xml", "_rels/.rels", "xl/workbook.xml",
                       "xl/styles.xml"} - set(zf.namelist())
        if missing:
            raise OSError(f"저장 결과에 필수 파트 누락: {sorted(missing)}")
        os.replace(tmp, path)
    except Exception as e:
        logger.warning("⚠️ 일일 리포트 저장 실패(기존 파일 보존): %s", e)
        try:
            os.remove(tmp)
        except OSError:
            pass
        return None

    logger.info("📊 일일 리포트 갱신: %s  (%s · hit %d · 실패/미실행 %d)", path, run_date, n_hit, n_fail)
    return path


# ─── 실행 ──────────────────────────────────────────────────────────

def run(sitecodes: list[str], pages: list[str], capture: bool, run_date: str) -> int:
    rules = parse_keyword_rules()
    if not rules:
        logger.error("❌ KEYWORDS 가 비어 있습니다 — 파일 상단에 한 줄 이상 적어 주세요.")
        return 2
    if not sitecodes or not pages:
        logger.error("❌ 대상 sitecode 또는 CHECK_PAGES 가 비어 있습니다.")
        return 2

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    csv_path = OUTPUT_DIR / RESULT_CSV_NAME
    try:
        csv_path.unlink()
    except FileNotFoundError:
        pass

    # 작업 단위 = (sitecode, page). sitecode 순서 안에서 CHECK_PAGES 순서.
    jobs = [(sc, pg) for sc in sitecodes for pg in pages]
    logger.info("🔑 키워드 규칙 %d개: %s", len(rules), " | ".join(line for line, _ in rules))
    logger.info("🌍 대상 %d개국 × 페이지 %s = %d건 · 캡처 %s · 리포트 날짜 %s",
                len(sitecodes), "/".join(pages), len(jobs), "ON" if capture else "OFF", run_date)

    # 캡처 기록 — 워커에는 읽기 전용 스냅샷을 주고, 갱신·저장은 메인 스레드에서만 한다.
    # (같은 (sitecode, page) 는 한 실행에서 한 번만 돌므로 스냅샷으로 충분하다)
    capture_state = load_capture_state()
    completed = False
    hits = []
    new_captures = 0
    # ⚠ ThreadPoolExecutor 를 쓰지 않는다 — Playwright sync 객체는 **만든 스레드에서만** 닫을 수 있어
    #   메인 스레드에서 세션을 정리하면 greenlet.error 로 실패하고 Chrome 이 남는다(2026-09-29 확인).
    #   워커 스레드가 큐를 비운 뒤 자기 세션을 직접 닫고, 결과는 result 큐로 메인 스레드에 넘긴다.
    job_q: queue.Queue = queue.Queue()
    for job in jobs:
        job_q.put(job)
    result_q: queue.Queue = queue.Queue()

    def _worker() -> None:
        try:
            while True:
                try:
                    sc, pg = job_q.get_nowait()
                except queue.Empty:
                    break
                try:
                    row = check_site_with_retry(sc, pg, rules, capture, run_date,
                                                capture_state.get(_state_key(sc, pg), {}))
                except Exception as e:          # 세션 기동 실패 등 — 결과는 반드시 한 줄 남긴다
                    row = _result(sc, pg, build_url(sc, pg), f"{RESULT_FAIL_PREFIX}(error)",
                                  detail=f"{type(e).__name__}: {str(e).splitlines()[0][:150]}")
                result_q.put(row)
        finally:
            _close_thread_session()

    workers = [threading.Thread(target=_worker, daemon=True) for _ in range(max(1, MAX_WORKERS))]
    try:
        for t in workers:
            t.start()
        for i in range(1, len(jobs) + 1):
            row = result_q.get()
            tag = f"{row['sitecode']}/{row['page']}"
            _append_result_csv(csv_path, row)     # 메인 스레드에서만 쓴다
            if row["new_elements"]:
                bucket = capture_state.setdefault(_state_key(row["sitecode"], row["page"]), {})
                for el in json.loads(row["new_elements"]):
                    bucket.setdefault(el["key"], {})[el["fp"]] = {
                        "first_seen": datetime.now().strftime("%Y-%m-%d"),
                        "kind": el["kind"], "capture": row["capture"], "design": el["design"],
                    }
                save_capture_state(capture_state)
                new_captures += 1
            if row["result"] == RESULT_HIT:
                hits.append(tag)
                snap = "📷 새로 캡처" if row["capture"] else "기존 캡처 있음 → 캡처 생략"
                logger.info("🎯 [%d/%d] %s: HIT — %s (%s)", i, len(jobs), tag, row["keywords"], snap)
            elif row["result"] == RESULT_MISS:
                logger.info("➖ [%d/%d] %s: 없음", i, len(jobs), tag)
            else:
                logger.info("❌ [%d/%d] %s: %s %s", i, len(jobs), tag, row["result"], row["detail"])
        completed = True
    finally:
        for t in workers:
            t.join(timeout=60)
        # 중간에 죽어도 그때까지의 결과는 리포트에 남긴다
        if csv_path.exists():
            write_daily_report(csv_path, run_date, jobs)
        if completed and not KEEP_RESULT_CSV:
            try:
                csv_path.unlink()
            except OSError:
                pass

    logger.info("✅ 완료 — HIT %d건 · 새 캡처 %d건%s", len(hits), new_captures,
                f": {', '.join(sorted(hits))}" if hits else "")
    return 0


class OperatorFormatter(logging.Formatter):
    def format(self, record: logging.LogRecord) -> str:
        return "" if record.getMessage() == "" else super().format(record)


def _setup_logging(debug: bool) -> None:
    handlers: list[logging.Handler] = []
    fmt = OperatorFormatter("[%(asctime)s] %(message)s", datefmt="%H:%M:%S")
    if sys.stderr is not None:              # pythonw 는 stderr 가 None
        try:
            sys.stderr.reconfigure(encoding="utf-8", errors="replace")   # cp949 콘솔에서 한글·이모지 깨짐 방지
        except Exception:
            pass
        h = logging.StreamHandler()
        h.setFormatter(fmt)
        handlers.append(h)
    if LOG_TO_FILE:
        OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
        fh = logging.FileHandler(OUTPUT_DIR / RUN_LOG_NAME, mode="w", encoding="utf-8")
        fh.setFormatter(OperatorFormatter("[%(asctime)s] %(message)s", datefmt="%Y-%m-%d %H:%M:%S"))
        handlers.append(fh)
    logging.basicConfig(level=logging.INFO, handlers=handlers)
    # root 를 DEBUG 로 내리면 playwright 로그가 섞인다 — 이 스크립트 logger 만 내린다.
    if debug:
        logger.setLevel(logging.DEBUG)


def main() -> None:
    parser = argparse.ArgumentParser(description="홈·프로모션(offer) 페이지에 타겟 캠페인 링크가 있는지 매일 확인합니다.")
    parser.add_argument("--sitecodes", default=None, help="쉼표 구분 sitecode (예: ae,hq). 미지정 시 SITECODES 전체")
    parser.add_argument("--pages", default=None, help="쉼표 구분 페이지 (예: home). 미지정 시 CHECK_PAGES")
    parser.add_argument("--no-capture", action="store_true", help="hit 여도 캡처하지 않음")
    parser.add_argument("--debug", action="store_true", help="진단 로그")
    parser.add_argument("--run-date", default=None, help=argparse.SUPPRESS)   # 리포트 날짜 블록 테스트용 (MMDD)
    args = parser.parse_args()

    _setup_logging(args.debug)
    all_sitecodes = parse_sitecodes()
    if args.sitecodes:
        requested = [s.strip().lower() for s in args.sitecodes.split(",") if s.strip()]
        unknown = [s for s in requested if s not in all_sitecodes]
        if unknown:
            logger.info("ℹ️ SITECODES 목록에 없는 코드도 그대로 실행합니다: %s", ", ".join(unknown))
        sitecodes = list(dict.fromkeys(requested))
    else:
        sitecodes = all_sitecodes

    started = time.monotonic()
    logger.info("🚀 promo_link_checker 시작")
    pages = parse_pages(args.pages.replace(",", "\n")) if args.pages else parse_pages()
    code = run(sitecodes, pages, CAPTURE_ON_HIT and not args.no_capture,
               args.run_date or datetime.now().strftime("%m%d"))
    logger.info("⏱️ 총 소요 %d초", round(time.monotonic() - started))
    raise SystemExit(code)


if __name__ == "__main__":
    main()
